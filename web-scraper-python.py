import requests
from bs4 import BeautifulSoup
import openpyxl

# 1. Definimos o endereço do site de testes
url = "http://quotes.toscrape.com/"

print(f"Conectando ao site: {url}...")
resposta = requests.get(url)

if resposta.status_code == 200:
    print("Conexão bem-sucedida! Raspando os dados...\n")

    soup = BeautifulSoup(resposta.text, 'html.parser')
    lista_citacoes = soup.find_all('div', class_='quote')

    # 2. Criamos uma nova planilha do Excel usando o openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Frases Famosas"  # Nome da aba na planilha

    # 3. Criamos o Cabeçalho Organizado (Linha 1)
    # Usaremos colunas separadas para a Frase e para o Autor
    ws.append(["Citação / Frase", "Autor"])

    # Estilizando o cabeçalho (Deixando em negrito para dar destaque profissional)
    for coluna in range(1, 3):
        celula = ws.cell(row=1, column=coluna)
        celula.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        celula.fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78",
                                                  fill_type="solid")  # Azul corporativo
        celula.alignment = openpyxl.styles.Alignment(horizontal="center", vertical="center")

    # 4. Percorremos as citações e adicionamos cada uma como uma linha na planilha
    contador = 0
    for item in lista_citacoes:
        texto = item.find('span', class_='text').text
        autor = item.find('small', class_='author').text

        # Adiciona a linha na planilha
        ws.append([texto, autor])

        contador += 1
        print(f"• Capturada [{contador}]: {autor}")

    # 5. Ajuste automático da largura das colunas para o texto não ficar cortado
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        # Define uma largura proporcional com um limite mínimo agradável
        ws.column_dimensions[col_letter].width = max(max_len + 5, 20)

    # 6. Salvamos a planilha com um nome limpo e profissional
    nome_arquivo = "Citacoes_Web_Scraping.xlsx"
    wb.save(nome_arquivo)

    print(f"\nSucesso absoluto! Dados salvos e formatados no arquivo: '{nome_arquivo}'")

else:
    print(f"Erro ao acessar a página. Código de status: {resposta.status_code}")